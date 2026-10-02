import io

from openpyxl import load_workbook


def _login(client):
    r = client.post("/api/auth/login", json={"username": "admin", "password": "Passw0rd!"})
    assert r.status_code == 200


def test_export_user_sheet_xlsx(client, seeded_user):
    _login(client)
    wb = client.get("/api/workbooks/default").json()
    user_sheet = next(s for s in wb["sheets"] if s["kind"] == "USER")

    client.patch(
        f"/api/sheets/{user_sheet['id']}/cells",
        json={"cells": [
            {"row": 1, "col": 1, "raw": "100"},
            {"row": 1, "col": 2, "raw": "200"},
        ]},
    )

    r = client.get(f"/api/exports/sheets/{user_sheet['id']}/xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb2 = load_workbook(io.BytesIO(r.content))
    ws = wb2.active
    assert ws.cell(row=1, column=1).value == "100"
    assert ws.cell(row=1, column=2).value == "200"


def test_export_system_sheet_rejected(client, seeded_user):
    _login(client)
    wb = client.get("/api/workbooks/default").json()
    ar = next(s for s in wb["sheets"] if s["name"] == "Asset Register")
    r = client.get(f"/api/exports/sheets/{ar['id']}/xlsx")
    assert r.status_code == 400


def test_export_asset_register_xlsx(client, seeded_user):
    _login(client)
    r = client.get("/api/exports/reports/asset-register/xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    # Title should be at A1
    assert ws.cell(row=1, column=1).value == "Asset Register"


def test_export_asset_register_pdf(client, seeded_user):
    _login(client)
    r = client.get("/api/exports/reports/asset-register/pdf")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    # PDF magic number
    assert r.content[:4] == b"%PDF"


def test_export_depreciation_schedule_pdf(client, seeded_user):
    _login(client)
    r = client.get("/api/exports/reports/depreciation-schedule/pdf")
    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"


def test_export_unknown_report_404(client, seeded_user):
    _login(client)
    r = client.get("/api/exports/reports/does-not-exist/xlsx")
    assert r.status_code == 404


def test_import_xlsx_creates_new_sheet(client, seeded_user):
    _login(client)
    wb = client.get("/api/workbooks/default").json()

    # Build a small xlsx in-memory
    from openpyxl import Workbook as XlWb
    xl = XlWb()
    ws = xl.active
    ws.cell(row=1, column=1, value="Hello")
    ws.cell(row=1, column=2, value=42)
    buf = io.BytesIO()
    xl.save(buf)
    buf.seek(0)

    r = client.post(
        "/api/imports/xlsx-to-sheet",
        data={"workbook_id": str(wb["id"]), "sheet_name": "Imported Data"},
        files={
            "file": (
                "test.xlsx",
                buf.read(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 201
    body = r.json()
    assert body["name"] == "Imported Data"
    assert body["kind"] == "USER"

    # Verify cell content
    render = client.get(f"/api/sheets/{body['id']}/render").json()
    cells = {(c["row"], c["col"]): c for c in render["grid"]["cells"]}
    assert cells[(1, 1)]["raw"] == "Hello"
    assert cells[(1, 2)]["raw"] == "42"


def test_import_requires_unique_name(client, seeded_user):
    _login(client)
    wb = client.get("/api/workbooks/default").json()

    from openpyxl import Workbook as XlWb
    xl = XlWb()
    xl.active.cell(row=1, column=1, value="x")
    buf = io.BytesIO()
    xl.save(buf)
    buf.seek(0)

    # First import succeeds
    r1 = client.post(
        "/api/imports/xlsx-to-sheet",
        data={"workbook_id": str(wb["id"]), "sheet_name": "Clash"},
        files={"file": ("a.xlsx", buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r1.status_code == 201

    # Second with same name fails
    buf.seek(0)
    r2 = client.post(
        "/api/imports/xlsx-to-sheet",
        data={"workbook_id": str(wb["id"]), "sheet_name": "Clash"},
        files={"file": ("b.xlsx", buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r2.status_code == 409


def test_import_rejects_oversized(client, seeded_user):
    _login(client)
    wb = client.get("/api/workbooks/default").json()

    big = b"0" * (6_000_000)
    r = client.post(
        "/api/imports/xlsx-to-sheet",
        data={"workbook_id": str(wb["id"]), "sheet_name": "TooBig"},
        files={"file": ("big.xlsx", big, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 413